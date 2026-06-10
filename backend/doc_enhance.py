"""Insert captured figures into an existing .docx or .tex document at the
detected placeholder locations, preserving the surrounding formatting and
continuing figure numbering.
"""

from __future__ import annotations

import io
import os
import re
from typing import Optional


def _insert_paragraph_after(paragraph):
    """Insert and return a new empty paragraph after the given one."""
    from docx.text.paragraph import Paragraph
    from docx.oxml.ns import qn
    new_p = paragraph._p.makeelement(qn('w:p'), {})
    paragraph._p.addnext(new_p)
    return Paragraph(new_p, paragraph._parent)


def enhance_docx(original_bytes: bytes, mappings: list[dict], output_path: str,
                 start_figure: int = 1) -> str:
    """mappings: [{placeholder, image_path, caption}].

    For each mapping, find the first paragraph containing `placeholder`, clear
    its text and insert the image + caption there. Figures are numbered
    sequentially starting at `start_figure`.
    """
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document(io.BytesIO(original_bytes))
    fig_no = start_figure
    used = set()

    for mapping in mappings:
        placeholder = (mapping.get("placeholder") or "").strip()
        image_path = mapping.get("image_path")
        caption = mapping.get("caption") or ""
        if not image_path or not os.path.exists(image_path):
            continue

        target = None
        for p in doc.paragraphs:
            if id(p._p) in used:
                continue
            if placeholder and placeholder[:40].lower() in p.text.lower():
                target = p
                break
        # Fall back to any paragraph that still looks like a placeholder token.
        if target is None:
            for p in doc.paragraphs:
                if id(p._p) in used:
                    continue
                if re.search(r"\[FIGURE\]|\{\{FIGURE|insert screenshot|todo image", p.text, re.I):
                    target = p
                    break
        if target is None:
            continue

        used.add(id(target._p))
        # Clear the placeholder paragraph text.
        for run in list(target.runs):
            run.text = ""
        # Add the image into the (now empty) placeholder paragraph.
        run = target.add_run()
        try:
            run.add_picture(image_path, width=Inches(6))
            target.alignment = WD_ALIGN_PARAGRAPH.CENTER
        except Exception:
            target.add_run(f"[Could not embed image: {image_path}]")
        # Caption paragraph after the image.
        cap_text = caption if caption.lower().startswith("figure") else f"Figure {fig_no} — {caption}" if caption else f"Figure {fig_no}"
        cap_p = _insert_paragraph_after(target)
        cap_run = cap_p.add_run(cap_text)
        cap_run.italic = True
        cap_run.font.size = Pt(10)
        cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fig_no += 1

    doc.save(output_path)
    return output_path


def enhance_tex(original_text: str, mappings: list[dict], output_path: str,
                start_figure: int = 1) -> str:
    """Replace placeholder tokens in a LaTeX document with figure environments.

    LaTeX auto-numbers figures, so numbering is preserved by insertion order.
    """
    text = original_text
    fig_no = start_figure

    def figure_block(image_path: str, caption: str) -> str:
        rel = os.path.basename(image_path)
        cap = caption or "Figure"
        return (
            "\\begin{figure}[H]\n\\centering\n"
            f"\\includegraphics[width=0.85\\textwidth]{{{rel}}}\n"
            f"\\caption{{{cap}}}\n\\end{{figure}}"
        )

    for mapping in mappings:
        placeholder = mapping.get("placeholder") or ""
        image_path = mapping.get("image_path")
        caption = mapping.get("caption") or ""
        if not image_path:
            continue
        block = figure_block(image_path, caption)
        replaced = False
        # Try matching the placeholder text first.
        if placeholder:
            snippet = placeholder.strip()[:40]
            if snippet and snippet in text:
                text = text.replace(snippet, block, 1)
                replaced = True
        # Otherwise replace a common placeholder token.
        if not replaced:
            new_text, n = re.subn(r"\[FIGURE\]|%\s*FIGURE\s*%|\\placeholder", block, text, count=1)
            if n:
                text = new_text
                replaced = True
        if replaced:
            fig_no += 1

    with open(output_path, "w", encoding="utf-8") as f:
        f.write(text)
    return output_path
