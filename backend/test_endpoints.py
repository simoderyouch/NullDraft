"""Live integration test for all NullDraft backend endpoints."""
import io
import os
import json
import httpx
from PIL import Image, ImageDraw

BASE = "http://127.0.0.1:8000"
TMP = "/tmp/nulldraft-test"
os.makedirs(TMP, exist_ok=True)


def make_screenshot(path, with_sensitive=False):
    img = Image.new("RGB", (800, 500), "white")
    d = ImageDraw.Draw(img)
    d.rectangle([20, 20, 780, 80], fill=(40, 80, 160))
    d.text((40, 40), "Login Page", fill="white")
    d.text((40, 140), "Username:", fill="black")
    d.rectangle([160, 135, 500, 165], outline="black")
    d.text((170, 140), "admin", fill="black")
    d.text((40, 200), "Password:", fill="black")
    d.rectangle([160, 195, 500, 225], outline="black")
    if with_sensitive:
        d.text((170, 200), "S3cr3tP@ss!", fill="black")
        d.text((40, 260), "Email: john.doe@example.com", fill="black")
        d.text((40, 300), "API Key: sk-1234567890abcdef", fill="black")
    d.rectangle([160, 340, 300, 380], fill=(40, 160, 80))
    d.text((180, 350), "Sign In", fill="white")
    img.save(path)
    return path


def img_bytes(path):
    with open(path, "rb") as f:
        return f.read()


results = {}


def check(name, ok, detail=""):
    results[name] = "PASS" if ok else "FAIL"
    print(f"[{'PASS' if ok else 'FAIL'}] {name} {detail}")


shot = make_screenshot(os.path.join(TMP, "step-01-login.png"), with_sensitive=True)

with httpx.Client(timeout=90) as c:
    # health
    r = c.get(f"{BASE}/health")
    check("health", r.status_code == 200 and r.json().get("ai_configured"))

    # analyze-instructions (txt)
    instr = b"1. Open the login page.\n2. Enter username and password.\n3. Click Sign In and verify the dashboard loads."
    r = c.post(f"{BASE}/analyze-instructions",
               files={"file": ("tp.txt", instr, "text/plain")})
    steps = r.json().get("steps", []) if r.status_code == 200 else []
    check("analyze-instructions", r.status_code == 200 and len(steps) >= 2,
          f"({len(steps)} steps)")

    # analyze-instructions (docx)
    from docx import Document
    docbuf = io.BytesIO()
    dd = Document()
    dd.add_paragraph("Step A: Install dependencies.")
    dd.add_paragraph("Step B: Run the server.")
    dd.save(docbuf)
    r = c.post(f"{BASE}/analyze-instructions",
               files={"file": ("tp.docx", docbuf.getvalue(),
                               "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    check("analyze-instructions(docx)", r.status_code == 200 and len(r.json().get("steps", [])) >= 1)

    # generate-description (vision)
    r = c.post(f"{BASE}/generate-description?step_title=Login%20page",
               files={"file": ("s.png", img_bytes(shot), "image/png")})
    desc = r.json().get("description", "") if r.status_code == 200 else ""
    check("generate-description", r.status_code == 200 and len(desc) > 10,
          f"({len(desc)} chars)")

    # validate-step (vision)
    r = c.post(f"{BASE}/validate-step?step_title=Login%20page&step_description=Show%20the%20login%20form",
               files={"file": ("s.png", img_bytes(shot), "image/png")})
    check("validate-step", r.status_code == 200 and "pass" in r.json(),
          str(r.json()) if r.status_code == 200 else r.text)

    # detect-sensitive (vision)
    r = c.post(f"{BASE}/detect-sensitive",
               files={"file": ("s.png", img_bytes(shot), "image/png")})
    regions = r.json().get("regions", []) if r.status_code == 200 else []
    check("detect-sensitive", r.status_code == 200, f"({len(regions)} regions)")

    # blur-regions (apply blur if any region; else use a dummy region)
    blur_in = shot
    blur_out = os.path.join(TMP, "step-01-blurred.png")
    use_regions = regions or [{"x": 0.2, "y": 0.38, "width": 0.4, "height": 0.08}]
    r = c.post(f"{BASE}/blur-regions",
               json={"input_path": blur_in, "regions": use_regions, "output_path": blur_out})
    check("blur-regions", r.status_code == 200 and os.path.exists(blur_out))

    # ocr (vision)
    r = c.post(f"{BASE}/ocr", files={"file": ("s.png", img_bytes(shot), "image/png")})
    text = r.json().get("text", "") if r.status_code == 200 else ""
    check("ocr", r.status_code == 200 and len(text) > 3, f"({len(text)} chars)")

    # optimize-image
    opt_out = os.path.join(TMP, "step-01-opt.jpg")
    r = c.post(f"{BASE}/optimize-image",
               json={"input_path": shot, "output_path": opt_out, "max_width": 400,
                     "quality": 80, "format": "jpg"})
    check("optimize-image", r.status_code == 200 and os.path.exists(opt_out))

    # generate-captions
    cap_steps = [{"id": "1", "title": "Login page", "description": "Show login",
                  "generated_description": desc}]
    r = c.post(f"{BASE}/generate-captions", json={"steps": cap_steps})
    caps = r.json().get("captions", []) if r.status_code == 200 else []
    check("generate-captions", r.status_code == 200 and len(caps) >= 1, str(caps))

    # generate-narrative
    r = c.post(f"{BASE}/generate-narrative",
               json={"title": "Login Guide",
                     "steps": [{"title": "Login page", "description": "Show login"}]})
    check("generate-narrative", r.status_code == 200 and "introduction" in r.json())

    # analyze-document
    docbuf2 = io.BytesIO()
    dd2 = Document()
    dd2.add_paragraph("Introduction.")
    dd2.add_paragraph("[FIGURE] Insert screenshot of the login page here.")
    dd2.add_paragraph("Conclusion.")
    dd2.save(docbuf2)
    r = c.post(f"{BASE}/analyze-document",
               files={"file": ("doc.docx", docbuf2.getvalue(),
                               "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    check("analyze-document", r.status_code == 200 and "placeholders" in r.json())

    # suggest-next
    r = c.post(f"{BASE}/suggest-next?current_title=Login%20page",
               files={"file": ("s.png", img_bytes(shot), "image/png")})
    check("suggest-next", r.status_code == 200 and "title" in r.json())

    # generate-report for all formats
    report_steps = [{
        "id": "1", "number": 1, "title": "Login page",
        "description": "Open the login page and enter credentials",
        "generated_description": desc or "The login page with username and password.",
        "generated_caption": "", "notes": "Use a test account.",
        "image_path": shot, "skipped": False,
    }, {
        "id": "2", "number": 2, "title": "Dashboard",
        "description": "Verify the dashboard loads",
        "generated_description": "", "generated_caption": "Main dashboard",
        "notes": "", "image_path": shot, "skipped": False,
    }]
    for fmt in ["md", "tex", "docx", "pdf", "json"]:
        out = os.path.join(TMP, f"report.{fmt}")
        r = c.post(f"{BASE}/generate-report", json={
            "format": fmt, "title": "Login Guide", "output_path": out,
            "steps": report_steps, "template": "academic",
            "branding": {"title": "Login Guide", "author": "Test", "subtitle": "QA Report"},
        })
        produced = r.json().get("output_path") if r.status_code == 200 else None
        ok = r.status_code == 200 and produced and os.path.exists(produced) and os.path.getsize(produced) > 0
        check(f"generate-report({fmt})", ok,
              f"-> {produced} ({os.path.getsize(produced) if produced and os.path.exists(produced) else 0} bytes)"
              if ok else r.text[:200])

    # back-compat latex endpoint
    r = c.post(f"{BASE}/generate-latex-report", json={
        "title": "Compat", "steps": [{"id": "1", "title": "A", "description": "B"}],
        "image_paths": {"1": shot}, "output_path": os.path.join(TMP, "compat"),
    })
    check("generate-latex-report", r.status_code == 200 and "tex_path" in r.json())

print("\n=== SUMMARY ===")
passed = sum(1 for v in results.values() if v == "PASS")
print(f"{passed}/{len(results)} passed")
for k, v in results.items():
    if v == "FAIL":
        print(f"  FAILED: {k}")
